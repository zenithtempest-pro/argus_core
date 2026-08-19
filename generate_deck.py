from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Initialize presentation with 16:9 widescreen layout
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
DARK_BG = RGBColor(9, 13, 22)
CARD_BG = RGBColor(19, 27, 46)
TEXT_WHITE = RGBColor(248, 250, 252)
TEXT_MUTED = RGBColor(148, 163, 184)
ACCENT_BLUE = RGBColor(37, 99, 235)
BORDER_COLOR = RGBColor(30, 41, 59)

slides_data = [
    {
        "title": "ArgusCore Intelligence & Diagnostics",
        "subtitle": "Unified OSINT, Network Diagnostics, and Automation Platform",
        "category": "PROJECT OVERVIEW",
        "bullets": [
            "Comprehensive passive reconnaissance across Identity, Network, Media, and Email.",
            "Integrated MXToolbox-grade diagnostic suite with real-time RBL and DNS validation.",
            "Centralized Script Hub for verified, standalone Python OSINT CLI tools.",
            "Engineered with clean Light/Dark UX, non-blocking async sockets, and RBAC admin controls."
        ]
    },
    {
        "title": "The Problem & Market Need",
        "subtitle": "Eliminating tool fragmentation in modern intelligence workflows",
        "category": "PROBLEM STATEMENT",
        "bullets": [
            "Fragmented Tooling: Investigators switch between 10+ disjointed sites for basic recon.",
            "UI Freezes & Timeouts: Synchronous socket calls frequently hang standard diagnostic sites.",
            "Complex Indian Telecom Parsing: Lack of unified DoT circle and carrier attribution for +91 numbers.",
            "Script Distribution Gap: Offline threat hunting scripts lack verified checksum repositories."
        ]
    },
    {
        "title": "ArgusCore System Architecture",
        "subtitle": "Modular, high-performance full-stack structure",
        "category": "TECHNICAL DESIGN",
        "bullets": [
            "Frontend: Next.js App Router, Tailwind CSS with dynamic dark/light theme persistence.",
            "Asynchronous Engine: Non-blocking parallel queries (Promise.allSettled) with 3.5s timeouts.",
            "Security Middleware: SSRF private IP blocking, input sanitization, and rate limiting.",
            "Modular Routing: Dedicated hubs for SuperTool, Downloads, Dashboards, and Administration."
        ]
    },
    {
        "title": "Passive Reconnaissance Suite",
        "subtitle": "Multi-dimensional target investigation",
        "category": "CORE OSINT",
        "bullets": [
            "Username Hunter: Asynchronous profile scanner probing 25+ major social networks.",
            "Indian (+91) Phone Engine: 4-digit TRAI/DoT prefix mapping across 22 telecom circles.",
            "Media Forensics: Metadata & GPS EXIF extractor and scrubber.",
            "Document Dorking: Automated search generation for PDF, DOCX, PPTX, and exposed indexes."
        ]
    },
    {
        "title": "MXToolbox-Grade Diagnostic Engine",
        "subtitle": "Deep email security and server reputation checks",
        "category": "NETWORK DIAGNOSTICS",
        "bullets": [
            "SuperTool Command Line: Universal syntax bar with prefix sanitization (mx:, blacklist:, etc.).",
            "Multi-RBL Blacklist Checker: Simultaneous IP lookup across 100+ DNSBLs (Spamhaus, Barracuda).",
            "Email Protocol Validation: Strict RFC 7489 DMARC, SPF, DKIM, and BIMI policy checks.",
            "Header Analyzer: Hop latency timeline, spam metrics, and transit analysis."
        ]
    },
    {
        "title": "Python Script Repository Hub",
        "subtitle": "Verified, standalone CLI tools for offline threat hunting",
        "category": "DEVELOPER ECOSYSTEM",
        "bullets": [
            "Dynamic Script Loading: Automatically detects and serves .py files added to /public/scripts/.",
            "Integrity & Metadata: Automated SHA-256 hash generation and dependency mapping.",
            "Curated Core Tools: Pre-loaded with DNS Recon Pro, Username Hunter, and EXIF Scrubber.",
            "Direct Download: Native one-click acquisition for analyst environments."
        ]
    },
    {
        "title": "Security, Administration & Support",
        "subtitle": "Role-based controls and operational safeguards",
        "category": "GOVERNANCE",
        "bullets": [
            "RBAC Admin Panel: Live user session monitoring, query logs, and account controls at /admin.",
            "Flippa-Style Authentication: Clean email login, standard auth, and Google/LinkedIn SSO.",
            "Official Support Ecosystem: Direct analyst routing to supportarguscore@gmail.com with SLA tracking.",
            "Defensive Protections: Command injection guards and CSP headers."
        ]
    }
]

for item in slides_data:
    slide = prs.slides.add_slide(blank_layout)
    
    # 1. Slide Background
    bg = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    bg.line.fill.background()
    
    # 2. Category Tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(6), Inches(0.4))
    tf_cat = cat_box.text_frame
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = item["category"]
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = ACCENT_BLUE
    
    # 3. Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.5), Inches(0.8))
    tf_t = title_box.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = item["title"]
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE
    
    # 4. Slide Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(11.5), Inches(0.5))
    tf_s = sub_box.text_frame
    p_s = tf_s.paragraphs[0]
    p_s.text = item["subtitle"]
    p_s.font.size = Pt(14)
    p_s.font.color.rgb = TEXT_MUTED
    
    # 5. Background Card
    card = slide.shapes.add_shape(1, Inches(0.8), Inches(2.3), Inches(11.733), Inches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COLOR
    
    # 6. Bullet Items Text Box
    content_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(10.9), Inches(3.8))
    tf_c = content_box.text_frame
    tf_c.word_wrap = True
    
    for i, bullet in enumerate(item["bullets"]):
        p = tf_c.paragraphs[0] if i == 0 else tf_c.add_paragraph()
        p.text = f"•   {bullet}"
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(18)

output_filename = "ArgusCore_Project_Presentation.pptx"
prs.save(output_filename)
print(f"Presentation successfully created and saved as '{output_filename}'")